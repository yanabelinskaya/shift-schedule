import csv
import io
import logging
import secrets

from django.conf import settings
from django.contrib.auth import login, logout
from django.contrib.auth import get_user_model
from django.core.mail import send_mail
from django.shortcuts import get_object_or_404
from drf_spectacular.utils import extend_schema
from rest_framework import generics, status
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.authtoken.models import Token
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from django.utils.decorators import method_decorator
from django.db.models import Count
from django.utils import timezone
from django.urls import reverse
from django.views.decorators.csrf import csrf_exempt

from .models import (
    Department,
    DepartmentPosition,
    EmployeeAbsence,
    EmployeeProfile,
    GlobalSettings,
    PasswordResetRequest,
)
from .object_access import can_manage_employee
from .pagination import BoundedListMixin, resolve_non_paginated_limit
from .permissions import IsAdminRole
from .serializers import (
    ApiRootSerializer,
    DepartmentArchiveSerializer,
    DepartmentCreateSerializer,
    DepartmentSerializer,
    DepartmentUpdateSerializer,
    DetailMessageSerializer,
    EmployeeActivateResponseSerializer,
    EmployeeAbsenceCreateSerializer,
    EmployeeAbsenceSerializer,
    EmployeeAvatarResponseSerializer,
    EmployeeAvatarUploadSerializer,
    EmployeeCreateSerializer,
    EmployeeCreateResponseSerializer,
    EmployeeDeactivateResponseSerializer,
    EmployeeIdsRequestSerializer,
    EmployeeImportRequestSerializer,
    EmployeeImportResponseSerializer,
    EmployeeSerializer,
    EmployeeUpdateSerializer,
    EmptySerializer,
    LoginSerializer,
    LoginResponseSerializer,
    PasswordResetPendingItemSerializer,
    PasswordResetRequestSerializer,
    PasswordResetResolveResponseSerializer,
    UserCreateSerializer,
    UserCreateResponseSerializer,
    UserSerializer,
)


logger = logging.getLogger(__name__)


def _platform_name():
    settings_obj = GlobalSettings.objects.only("platform_name").first()
    return (settings_obj.platform_name if settings_obj else "") or "Conector Shift"


def _send_credentials_email(user, raw_password):
    if not user.email or not raw_password:
        return False, "Email или пароль отсутствуют."
    platform_name = _platform_name()
    subject = f'{platform_name} — доступ к аккаунту'
    message = (
        f'Здравствуйте, {user.get_full_name() or user.username}!\n\n'
        f'Для вас создан аккаунт в сервисе {platform_name}.\n\n'
        f'Логин: {user.username}\n'
        f'Временный пароль: {raw_password}\n\n'
        'Рекомендуем изменить пароль после первого входа.\n\n'
        'Если вы не ожидали это письмо, просто проигнорируйте его.'
    )
    try:
        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [user.email],
            fail_silently=False,
        )
    except Exception as exc:
        logger.exception("Failed to send credentials email to %s", user.email)
        return False, str(exc)
    return True, ""


def _send_reset_email(user, raw_password):
    if not user.email or not raw_password:
        return False, "Email или пароль отсутствуют."
    platform_name = _platform_name()
    subject = f'{platform_name} — восстановление доступа'
    message = (
        f'Здравствуйте, {user.get_full_name() or user.username}!\n\n'
        f'Вы запросили восстановление доступа к {platform_name}.\n\n'
        f'Логин: {user.username}\n'
        f'Новый пароль: {raw_password}\n\n'
        'Рекомендуем изменить пароль после входа.\n\n'
        'Если вы не запрашивали восстановление, сообщите администратору.'
    )
    try:
        send_mail(
            subject,
            message,
            settings.DEFAULT_FROM_EMAIL,
            [user.email],
            fail_silently=False,
        )
    except Exception as exc:
        logger.exception("Failed to send reset email to %s", user.email)
        return False, str(exc)
    return True, ""


def _normalize_header(value):
    return " ".join(str(value or "").strip().lower().split())


def _map_role(value):
    normalized = _normalize_header(value)
    if normalized in {"администратор", "admin"}:
        return "admin"
    if normalized in {"менеджер", "manager"}:
        return "manager"
    if normalized in {"сотрудник", "employee", "staff"}:
        return "employee"
    return ""


def _row_to_payload(headers, row):
    header_map = {
        "фио": "full_name",
        "full_name": "full_name",
        "full name": "full_name",
        "name": "full_name",
        "фамилия": "last_name",
        "lastname": "last_name",
        "last name": "last_name",
        "имя": "first_name",
        "firstname": "first_name",
        "first name": "first_name",
        "отчество": "middle_name",
        "middlename": "middle_name",
        "middle name": "middle_name",
        "email": "email",
        "e-mail": "email",
        "почта": "email",
        "телефон": "corporate_phone",
        "phone": "corporate_phone",
        "роль": "role",
        "role": "role",
        "отдел": "department_name",
        "department": "department_name",
        "должность": "position",
        "position": "position",
        "ставка": "monthly_salary",
        "оклад": "monthly_salary",
        "salary": "monthly_salary",
        "rate": "monthly_salary",
    }
    payload = {}
    for index, header in enumerate(headers):
        key = header_map.get(header)
        if not key:
            continue
        value = row[index] if index < len(row) else ""
        if value is None:
            value = ""
        value = str(value).strip()
        if key == "role":
            value = _map_role(value)
        if key == "monthly_salary" and not value:
            continue
        payload[key] = value
    return payload


def _parse_csv(file_obj):
    content = file_obj.read()
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1251", errors="ignore")
    stream = io.StringIO(text)
    sample = text[:1024]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";"
    reader = csv.reader(stream, dialect)
    rows = list(reader)
    if not rows:
        return []
    headers = [_normalize_header(cell) for cell in rows[0]]
    return [(_row_to_payload(headers, row), index + 2) for index, row in enumerate(rows[1:])]


def _parse_xlsx(file_obj):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("Установите openpyxl для импорта XLSX.") from exc

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_normalize_header(cell) for cell in rows[0]]
    return [(_row_to_payload(headers, list(row)), index + 2) for index, row in enumerate(rows[1:])]


def _parse_xls(file_obj):
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("Установите xlrd для импорта XLS.") from exc

    workbook = xlrd.open_workbook(file_contents=file_obj.read())
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return []
    headers = [_normalize_header(cell) for cell in sheet.row_values(0)]
    payloads = []
    for index in range(1, sheet.nrows):
        payloads.append((_row_to_payload(headers, sheet.row_values(index)), index + 1))
    return payloads


def _format_serializer_errors(errors):
    field_labels = {
        "full_name": "ФИО",
        "first_name": "Имя",
        "last_name": "Фамилия",
        "middle_name": "Отчество",
        "email": "Email",
        "corporate_phone": "Телефон",
        "department_id": "Отдел",
        "department_name": "Отдел",
        "position": "Должность",
        "role": "Роль",
        "username": "Логин",
        "monthly_salary": "Оклад",
        "salary_reason": "Причина изменения оклада",
        "non_field_errors": "Ошибка",
    }
    parts = []
    for key, value in errors.items():
        if isinstance(value, (list, tuple)):
            message = " ".join(str(item) for item in value)
        else:
            message = str(value)
        label = field_labels.get(key, key)
        parts.append(f"{label}: {message}")
    return " ".join(parts)


class ApiRootView(generics.GenericAPIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    serializer_class = ApiRootSerializer

    @extend_schema(
        tags=['System'],
        summary='API root',
        description='Навигационная точка для основных API-эндпоинтов и документации.',
        responses={200: ApiRootSerializer},
    )
    def get(self, request):
        base = request.build_absolute_uri
        return Response(
            {
                'schema': base(reverse('schema')),
                'swagger': base(reverse('swagger-ui')),
                'redoc': base(reverse('redoc')),
                'auth_login': base(reverse('api-login')),
                'auth_me': base(reverse('api-me')),
                'users': base(reverse('api-user-create')),
                'departments': base(reverse('api-departments')),
                'employees': base(reverse('api-employees')),
                'password_resets': base(reverse('api-password-reset-request')),
                'admin_settings': base(reverse('api-admin-settings')),
                'admin_settings_history': base(reverse('api-admin-settings-history')),
                'admin_system_backups': base(reverse('api-admin-system-backups-create')),
                'admin_system_monitoring': base(reverse('api-admin-system-monitoring')),
                'manager_tasks': base(reverse('api-manager-tasks')),
                'manager_shift_requests': base(reverse('api-manager-shift-requests')),
                'employee_tasks': base(reverse('api-employee-tasks')),
                'employee_availability': base(reverse('api-employee-availability')),
            }
        )


@method_decorator(csrf_exempt, name='dispatch')
class LoginView(generics.GenericAPIView):
    permission_classes = [AllowAny]
    authentication_classes = []
    serializer_class = LoginSerializer

    @extend_schema(
        tags=['Authentication'],
        summary='Авторизация пользователя',
        description='Проверяет логин/пароль, создаёт токен DRF и возвращает данные текущего пользователя.',
        request=LoginSerializer,
        responses={
            200: LoginResponseSerializer,
            400: DetailMessageSerializer,
        },
    )
    def post(self, request):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        user = serializer.validated_data['user']

        login(request, user)
        if not request.session.session_key:
            request.session.save()
        token, _ = Token.objects.get_or_create(user=user)

        return Response(
            {
                'token': token.key,
                'user': UserSerializer(user).data,
                'session_key': request.session.session_key,
            }
        )


@method_decorator(csrf_exempt, name='dispatch')
class LogoutView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = EmptySerializer

    @extend_schema(
        tags=['Authentication'],
        summary='Выход из системы',
        description='Удаляет текущий токен пользователя и завершает сессию.',
        request=None,
        responses={204: None},
    )
    def post(self, request):
        token = getattr(request.user, 'auth_token', None)
        if token:
            token.delete()
        logout(request)
        return Response(status=status.HTTP_204_NO_CONTENT)


class MeView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = UserSerializer

    @extend_schema(
        tags=['Authentication'],
        summary='Текущий пользователь',
        description='Возвращает профиль авторизованного пользователя.',
        responses={200: UserSerializer},
    )
    def get(self, request):
        return Response(UserSerializer(request.user).data)


class UserCreateView(generics.CreateAPIView):
    serializer_class = UserCreateSerializer
    permission_classes = [IsAdminRole]

    def perform_create(self, serializer):
        user = serializer.save()
        EmployeeProfile.objects.get_or_create(user=user)
        raw_password = getattr(serializer, 'generated_password', None)

        email_sent, email_error = _send_credentials_email(user, raw_password)
        serializer.email_sent = email_sent
        serializer.email_error = email_error

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        headers = self.get_success_headers(serializer.data)
        raw_password = getattr(serializer, 'generated_password', None)
        return Response(
            {
                'user': serializer.data,
                'temporary_password': raw_password,
                'password_sent': bool(getattr(serializer, 'email_sent', False)),
                'email_error': getattr(serializer, 'email_error', ''),
            },
            status=status.HTTP_201_CREATED,
            headers=headers,
        )

    @extend_schema(
        tags=['Users'],
        summary='Создать пользователя',
        description='Создаёт пользователя (обычно менеджера/сотрудника) и отправляет временный пароль по email.',
        request=UserCreateSerializer,
        responses={
            201: UserCreateResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request, *args, **kwargs):
        return self.create(request, *args, **kwargs)


class DepartmentListView(BoundedListMixin, generics.ListCreateAPIView):
    permission_classes = [IsAdminRole]

    def get_queryset(self):
        return (
            Department.objects.select_related('manager')
            .prefetch_related('positions')
            .annotate(
                employees_count=Count('employees', distinct=True),
                positions_count=Count('positions', distinct=True),
            )
            .order_by('name')
        )

    def get_serializer_class(self):
        if self.request.method == 'POST':
            return DepartmentCreateSerializer
        return DepartmentSerializer

    @extend_schema(
        tags=['Departments'],
        summary='Список отделов',
        description='Возвращает список отделов с менеджерами, позициями и агрегированными счетчиками.',
        responses={200: DepartmentSerializer(many=True)},
    )
    def get(self, request, *args, **kwargs):
        return self.list(request, *args, **kwargs)

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        department = serializer.save()
        data = DepartmentSerializer(department, context=self.get_serializer_context()).data
        return Response(data, status=status.HTTP_201_CREATED)

    @extend_schema(
        tags=['Departments'],
        summary='Создать отдел',
        description='Создаёт отдел, связывает менеджера и список должностей.',
        request=DepartmentCreateSerializer,
        responses={
            201: DepartmentSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request, *args, **kwargs):
        return self.create(request, *args, **kwargs)


class DepartmentArchiveView(generics.UpdateAPIView):
    serializer_class = DepartmentArchiveSerializer
    permission_classes = [IsAdminRole]
    lookup_url_kwarg = 'department_id'
    http_method_names = ['patch']

    def get_queryset(self):
        return Department.objects.all()

    @extend_schema(
        tags=['Departments'],
        summary='Архивировать/разархивировать отдел',
        description='Обновляет флаг `is_archived` для отдела.',
        request=DepartmentArchiveSerializer,
        responses={
            200: DepartmentSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def patch(self, request, *args, **kwargs):
        return self.update(request, *args, **kwargs)

    def update(self, request, *args, **kwargs):
        partial = kwargs.pop('partial', True)
        instance = self.get_object()
        serializer = self.get_serializer(instance, data=request.data, partial=partial)
        serializer.is_valid(raise_exception=True)
        serializer.save()
        data = DepartmentSerializer(instance, context=self.get_serializer_context()).data
        return Response(data)


class DepartmentDetailView(generics.GenericAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = DepartmentUpdateSerializer

    @extend_schema(
        tags=['Departments'],
        summary='Обновить отдел',
        description='Обновляет название, менеджера и должности отдела.',
        request=DepartmentUpdateSerializer,
        responses={
            200: DepartmentSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def patch(self, request, department_id):
        department = get_object_or_404(Department, id=department_id)
        if department.is_archived:
            return Response(
                {'detail': 'Архивный отдел нельзя редактировать.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        serializer = self.get_serializer(
            data=request.data,
            context={'department': department},
            partial=True,
        )
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        update_fields = []
        if 'name' in data:
            department.name = data['name']
            update_fields.append('name')
        if 'manager_id' in data:
            manager_id = data.get('manager_id')
            manager = (
                get_user_model().objects.filter(id=manager_id, role='manager', is_active=True).first()
                if manager_id
                else None
            )
            department.manager = manager
            update_fields.append('manager')

        if update_fields:
            department.save(update_fields=update_fields)

        if 'positions' in data:
            desired = set(data.get('positions', []))
            existing = set(
                DepartmentPosition.objects.filter(department=department).values_list('title', flat=True)
            )
            to_add = desired - existing
            to_remove = existing - desired
            if to_remove:
                DepartmentPosition.objects.filter(
                    department=department,
                    title__in=to_remove,
                ).delete()
            if to_add:
                DepartmentPosition.objects.bulk_create(
                    [DepartmentPosition(department=department, title=title) for title in sorted(to_add)]
                )

        return Response(DepartmentSerializer(department).data, status=status.HTTP_200_OK)


class EmployeeListCreateView(BoundedListMixin, generics.ListCreateAPIView):
    permission_classes = [IsAdminRole]

    def get_queryset(self):
        from django.contrib.auth import get_user_model

        User = get_user_model()
        return (
            User.objects.exclude(role='admin')
            .select_related('profile__department')
            .order_by('last_name', 'first_name')
        )

    def get_serializer_class(self):
        if self.request.method == 'POST':
            return EmployeeCreateSerializer
        return EmployeeSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Список сотрудников',
        description='Возвращает сотрудников (кроме администраторов) с расширенными полями профиля.',
        responses={200: EmployeeSerializer(many=True)},
    )
    def get(self, request, *args, **kwargs):
        return self.list(request, *args, **kwargs)

    def perform_create(self, serializer):
        user = serializer.save()
        raw_password = getattr(serializer, 'generated_password', None)

        email_sent, email_error = _send_credentials_email(user, raw_password)
        serializer.email_sent = email_sent
        serializer.email_error = email_error

    def create(self, request, *args, **kwargs):
        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        self.perform_create(serializer)
        user = serializer.instance
        headers = self.get_success_headers({})
        return Response(
            {
                'user': EmployeeSerializer(user).data,
                'temporary_password': getattr(serializer, 'generated_password', None),
                'password_sent': bool(getattr(serializer, 'email_sent', False)),
                'email_error': getattr(serializer, 'email_error', ''),
            },
            status=status.HTTP_201_CREATED,
            headers=headers,
        )

    @extend_schema(
        tags=['Employees'],
        summary='Создать сотрудника',
        description='Создаёт сотрудника/менеджера, профиль и возвращает временный пароль.',
        request=EmployeeCreateSerializer,
        responses={
            201: EmployeeCreateResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request, *args, **kwargs):
        return self.create(request, *args, **kwargs)


class EmployeeDetailView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = EmployeeUpdateSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Обновить сотрудника',
        description='Обновляет данные сотрудника. Админ может менять все поля, сотрудник — только личные.',
        request=EmployeeUpdateSerializer,
        responses={
            200: EmployeeSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def patch(self, request, user_id):
        User = get_user_model()
        user = get_object_or_404(User, id=user_id)
        is_admin = getattr(request.user, 'role', None) == 'admin'
        if not is_admin and request.user.id != user.id:
            return Response({'detail': 'Недостаточно прав.'}, status=status.HTTP_403_FORBIDDEN)
        if user.role == 'admin' and not is_admin:
            return Response({'detail': 'Недостаточно прав.'}, status=status.HTTP_403_FORBIDDEN)
        profile, _ = EmployeeProfile.objects.get_or_create(user=user)

        payload = request.data.copy()
        if not is_admin:
            allowed_fields = {
                'first_name',
                'last_name',
                'middle_name',
                'corporate_phone',
                'personal_phone',
                'address',
                'passport_series',
                'passport_number',
                'passport_issued_by',
                'passport_issue_date',
                'snils',
                'inn',
            }
            payload = {key: value for key, value in payload.items() if key in allowed_fields}

        serializer = self.get_serializer(data=payload, context={'user': user})
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        user_fields = []
        if is_admin and 'first_name' in data:
            user.first_name = data.get('first_name', '')
            user_fields.append('first_name')
        if is_admin and 'last_name' in data:
            user.last_name = data.get('last_name', '')
            user_fields.append('last_name')
        if is_admin and 'email' in data:
            user.email = data.get('email', '')
            user_fields.append('email')
        if is_admin and 'is_active' in data:
            user.is_active = data['is_active']
            user_fields.append('is_active')
        if user_fields:
            user.save(update_fields=user_fields)
            if is_admin and user.role == 'manager' and 'is_active' in data and not user.is_active:
                Department.objects.filter(manager=user).update(manager=None)

        profile_fields = []

        def set_profile(field, value):
            setattr(profile, field, value)
            profile_fields.append(field)

        if is_admin and 'department_id' in data:
            department_id = data.get('department_id')
            department = Department.objects.filter(id=department_id).first() if department_id else None
            profile.department = department
            profile_fields.append('department')
        if is_admin and 'position' in data:
            set_profile('position', data.get('position', ''))
        if is_admin and 'middle_name' in data:
            set_profile('middle_name', data.get('middle_name', ''))
        if 'corporate_phone' in data:
            set_profile('corporate_phone', data.get('corporate_phone', ''))
        if 'personal_phone' in data:
            set_profile('personal_phone', data.get('personal_phone', ''))
        if 'address' in data:
            set_profile('address', data.get('address', ''))
        if is_admin and 'monthly_salary' in data:
            set_profile('monthly_salary', data.get('monthly_salary'))
        if is_admin and 'salary_reason' in data:
            set_profile('salary_reason', data.get('salary_reason', ''))
        if is_admin and 'department_change_reason' in data:
            set_profile('department_change_reason', data.get('department_change_reason', ''))
        if 'passport_series' in data:
            set_profile('passport_series', data.get('passport_series', ''))
        if 'passport_number' in data:
            set_profile('passport_number', data.get('passport_number', ''))
        if 'passport_issued_by' in data:
            set_profile('passport_issued_by', data.get('passport_issued_by', ''))
        if 'passport_issue_date' in data:
            set_profile('passport_issue_date', data.get('passport_issue_date'))
        if 'snils' in data:
            set_profile('snils', data.get('snils', ''))
        if 'inn' in data:
            set_profile('inn', data.get('inn', ''))

        if profile_fields:
            profile.save(update_fields=profile_fields)

        user = get_user_model().objects.select_related('profile__department').get(id=user.id)
        return Response(EmployeeSerializer(user).data, status=status.HTTP_200_OK)


class EmployeeAbsenceView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = EmployeeAbsenceCreateSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Список отсутствий сотрудника',
        description='Возвращает отпуска/больничные сотрудника для админа или менеджера его отдела.',
        responses={
            200: EmployeeAbsenceSerializer(many=True),
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def get(self, request, user_id):
        User = get_user_model()
        user = get_object_or_404(User, id=user_id)
        if user.role != 'employee' or not can_manage_employee(request.user, user):
            return Response({'detail': 'Недостаточно прав.'}, status=status.HTTP_403_FORBIDDEN)
        absences = EmployeeAbsence.objects.filter(user=user).order_by('-start_date', '-end_date')
        return Response(EmployeeAbsenceSerializer(absences, many=True).data, status=status.HTTP_200_OK)

    @extend_schema(
        tags=['Employees'],
        summary='Создать отсутствие сотрудника',
        description='Добавляет отпуск или больничный сотруднику с проверкой пересечений.',
        request=EmployeeAbsenceCreateSerializer,
        responses={
            201: EmployeeAbsenceSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def post(self, request, user_id):
        User = get_user_model()
        user = get_object_or_404(User, id=user_id)
        if user.role != 'employee' or not can_manage_employee(request.user, user):
            return Response({'detail': 'Недостаточно прав.'}, status=status.HTTP_403_FORBIDDEN)
        serializer = self.get_serializer(data=request.data, context={'user': user})
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data
        absence = EmployeeAbsence.objects.create(
            user=user,
            absence_type=data['absence_type'],
            start_date=data['start_date'],
            end_date=data['end_date'],
            created_by=request.user,
        )
        return Response(EmployeeAbsenceSerializer(absence).data, status=status.HTTP_201_CREATED)


class EmployeeImportView(generics.GenericAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = EmployeeImportRequestSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Импорт сотрудников из файла',
        description='Импортирует сотрудников из CSV/XLS/XLSX и возвращает список созданных записей и ошибок.',
        request=EmployeeImportRequestSerializer,
        responses={
            200: EmployeeImportResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request):
        file_obj = request.FILES.get('file') or request.FILES.get('employees_file')
        if not file_obj:
            return Response({'detail': 'Файл не выбран.'}, status=status.HTTP_400_BAD_REQUEST)

        extension = file_obj.name.split('.')[-1].lower()
        try:
            if extension == 'csv':
                payloads = _parse_csv(file_obj)
            elif extension == 'xlsx':
                payloads = _parse_xlsx(file_obj)
            elif extension == 'xls':
                payloads = _parse_xls(file_obj)
            else:
                return Response(
                    {'detail': 'Поддерживаются файлы XLSX, XLS, CSV.'},
                    status=status.HTTP_400_BAD_REQUEST,
                )
        except ImportError as exc:
            return Response({'detail': str(exc)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            return Response({'detail': 'Не удалось прочитать файл.'}, status=status.HTTP_400_BAD_REQUEST)

        created = []
        errors = []
        passwords_sent = 0
        for payload, row_number in payloads:
            if not any(value for value in payload.values()):
                continue
            serializer = EmployeeCreateSerializer(data=payload)
            if serializer.is_valid():
                user = serializer.save()
                raw_password = getattr(serializer, 'generated_password', None)
                email_sent, _ = _send_credentials_email(user, raw_password)
                if email_sent:
                    passwords_sent += 1
                created.append(EmployeeSerializer(user).data)
            else:
                errors.append(f"Строка {row_number}: {_format_serializer_errors(serializer.errors)}")

        return Response(
            {
                'created': created,
                'errors': errors,
                'created_count': len(created),
                'error_count': len(errors),
                'passwords_sent': passwords_sent,
            },
            status=status.HTTP_200_OK,
        )


class EmployeeAvatarView(generics.GenericAPIView):
    permission_classes = [IsAuthenticated]
    serializer_class = EmployeeAvatarUploadSerializer
    parser_classes = [MultiPartParser, FormParser]

    @extend_schema(
        tags=['Employees'],
        summary='Загрузить аватар сотрудника',
        description='Заменяет аватар текущего пользователя.',
        request=EmployeeAvatarUploadSerializer,
        responses={
            200: EmployeeAvatarResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def post(self, request, user_id):
        from django.contrib.auth import get_user_model

        User = get_user_model()
        if request.user.id != user_id:
            return Response({'detail': 'Недостаточно прав.'}, status=status.HTTP_403_FORBIDDEN)
        user = get_object_or_404(User, id=user_id)
        file_obj = request.FILES.get('avatar')
        if not file_obj:
            return Response({'detail': 'Файл не выбран.'}, status=status.HTTP_400_BAD_REQUEST)

        profile, _ = EmployeeProfile.objects.get_or_create(user=user)
        profile.avatar = file_obj
        profile.save(update_fields=['avatar'])
        avatar_url = request.build_absolute_uri(profile.avatar.url)
        return Response({'avatar_url': avatar_url}, status=status.HTTP_200_OK)


class PasswordResetRequestView(generics.GenericAPIView):
    permission_classes = [AllowAny]
    serializer_class = PasswordResetRequestSerializer

    def get_permissions(self):
        if self.request.method == 'GET':
            return [IsAdminRole()]
        return [AllowAny()]

    @extend_schema(
        tags=['Password Resets'],
        summary='Список заявок на сброс пароля',
        description='Возвращает заявки со статусом `pending` для администратора.',
        request=None,
        responses={
            200: PasswordResetPendingItemSerializer(many=True),
            403: DetailMessageSerializer,
        },
    )
    def get(self, request):
        pending = (
            PasswordResetRequest.objects.filter(status='pending')
            .select_related('user')
            .order_by('-created_at')
        )
        if request.query_params.get('page') or request.query_params.get('page_size'):
            page = self.paginate_queryset(pending)
            data = [
                {
                    'id': item.id,
                    'user_id': item.user_id,
                    'email': item.email,
                    'full_name': item.user.get_full_name() or item.user.username,
                    'created_at': item.created_at,
                }
                for item in page
            ]
            return self.get_paginated_response(data)

        limit = resolve_non_paginated_limit(request)
        data = [
            {
                'id': item.id,
                'user_id': item.user_id,
                'email': item.email,
                'full_name': item.user.get_full_name() or item.user.username,
                'created_at': item.created_at,
            }
            for item in pending[:limit]
        ]
        return Response(data, status=status.HTTP_200_OK)

    @extend_schema(
        tags=['Password Resets'],
        summary='Создать заявку на сброс пароля',
        description='Создаёт заявку на сброс пароля по корпоративному email.',
        request=PasswordResetRequestSerializer,
        responses={
            200: DetailMessageSerializer,
            400: DetailMessageSerializer,
        },
    )
    def post(self, request):
        settings_obj = GlobalSettings.objects.only("allow_password_reset_requests").first()
        if settings_obj is not None and not settings_obj.allow_password_reset_requests:
            return Response(
                {'detail': 'Восстановление пароля временно отключено администратором.'},
                status=status.HTTP_403_FORBIDDEN,
            )

        email = request.data.get('email', '').strip()
        if not email:
            return Response({'detail': 'Укажите корпоративную почту.'}, status=status.HTTP_400_BAD_REQUEST)
        User = get_user_model()
        user = User.objects.filter(email__iexact=email).exclude(role='admin').first()
        if not user:
            return Response({'detail': 'Сотрудник с такой почтой не найден.'}, status=status.HTTP_400_BAD_REQUEST)
        if not user.is_active:
            return Response({'detail': 'Аккаунт деактивирован. Обратитесь к администратору.'}, status=status.HTTP_400_BAD_REQUEST)

        existing = PasswordResetRequest.objects.filter(user=user, status='pending').first()
        if existing:
            return Response({'detail': 'Запрос уже отправлен и ожидает администратора.'}, status=status.HTTP_200_OK)

        PasswordResetRequest.objects.create(user=user, email=email)
        return Response({'detail': 'Запрос отправлен администратору.'}, status=status.HTTP_200_OK)


class PasswordResetResolveView(generics.GenericAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = EmptySerializer

    @extend_schema(
        tags=['Password Resets'],
        summary='Обработать заявку на сброс',
        description='Генерирует новый пароль по заявке, отправляет его сотруднику и закрывает заявку.',
        request=None,
        responses={
            200: PasswordResetResolveResponseSerializer,
            403: DetailMessageSerializer,
            404: DetailMessageSerializer,
        },
    )
    def post(self, request, request_id):
        reset_request = get_object_or_404(
            PasswordResetRequest,
            id=request_id,
            status='pending',
        )
        user = reset_request.user
        raw_password = secrets.token_urlsafe(8)
        user.set_password(raw_password)
        user.save(update_fields=['password'])

        email_sent, email_error = _send_reset_email(user, raw_password)
        reset_request.status = 'processed'
        reset_request.processed_at = timezone.now()
        reset_request.save(update_fields=['status', 'processed_at'])
        return Response(
            {
                'user_id': user.id,
                'email': user.email,
                'full_name': user.get_full_name() or user.username,
                'password_sent': email_sent,
                'email_error': email_error,
            },
            status=status.HTTP_200_OK,
        )


class EmployeeDeactivateView(generics.GenericAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = EmployeeIdsRequestSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Деактивировать сотрудников',
        description='Деактивирует список сотрудников и снимает менеджеров с закреплённых отделов.',
        request=EmployeeIdsRequestSerializer,
        responses={
            200: EmployeeDeactivateResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request):
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'detail': 'Выберите сотрудников.'}, status=status.HTTP_400_BAD_REQUEST)
        User = get_user_model()
        users = User.objects.filter(id__in=ids).exclude(role='admin')
        updated_ids = list(users.values_list('id', flat=True))
        managers = users.filter(role='manager')
        manager_departments = []
        if managers.exists():
            departments = (
                Department.objects.filter(manager__in=managers)
                .select_related('manager')
                .order_by('name')
            )
            manager_map = {}
            for department in departments:
                manager = department.manager
                if not manager:
                    continue
                entry = manager_map.setdefault(
                    manager.id,
                    {
                        'manager_id': manager.id,
                        'manager_name': manager.get_full_name() or manager.username,
                        'departments': [],
                    },
                )
                entry['departments'].append(
                    {
                        'id': department.id,
                        'name': department.name,
                        'is_archived': department.is_archived,
                        'detail_url': reverse('admin-department-detail', args=[department.id]),
                    }
                )
            manager_departments = list(manager_map.values())
            Department.objects.filter(manager__in=managers).update(manager=None)
        users.update(is_active=False)
        return Response(
            {
                'updated_ids': updated_ids,
                'manager_departments': manager_departments,
            },
            status=status.HTTP_200_OK,
        )


class EmployeeActivateView(generics.GenericAPIView):
    permission_classes = [IsAdminRole]
    serializer_class = EmployeeIdsRequestSerializer

    @extend_schema(
        tags=['Employees'],
        summary='Активировать сотрудников',
        description='Активирует список сотрудников без изменения текущих паролей.',
        request=EmployeeIdsRequestSerializer,
        responses={
            200: EmployeeActivateResponseSerializer,
            400: DetailMessageSerializer,
            403: DetailMessageSerializer,
        },
    )
    def post(self, request):
        ids = request.data.get('ids', [])
        if not isinstance(ids, list) or not ids:
            return Response({'detail': 'Выберите сотрудников.'}, status=status.HTTP_400_BAD_REQUEST)
        User = get_user_model()
        users = User.objects.filter(id__in=ids).exclude(role='admin')
        updated_ids = list(users.values_list('id', flat=True))
        users.update(is_active=True)
        return Response(
            {
                'updated_ids': updated_ids,
                'password_unchanged': True,
            },
            status=status.HTTP_200_OK,
        )
