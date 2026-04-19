import calendar
import csv
import io
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from django.contrib.auth import get_user_model
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db.models import Avg, Count, Max, Prefetch, Q, Sum
from django.db.models.functions import TruncDate
from django.http import FileResponse, Http404, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import ensure_csrf_cookie
from django.views.decorators.http import require_http_methods

from .models import (
    Department,
    DepartmentPosition,
    DepartmentTask,
    EmployeeProfile,
    GlobalSettings,
    PasswordResetRequest,
    SystemBackup,
    SystemLogEntry,
    TaskSubmission,
)
from .system_utils import collect_runtime_metrics, format_bytes, maybe_create_daily_backup
from .web_views_shared import (
    _ensure_role,
    _get_department_positions,
    _render_admin_page,
    _resolve_back_url,
    _serialize_log_entry,
    _system_monitoring_snapshot,
)

@login_required
def admin_dashboard(request):
    return admin_users(request)


@login_required
def admin_users(request):
    _ensure_role(request, 'admin')
    User = get_user_model()
    users = list(
        User.objects.exclude(role='admin')
        .order_by('last_name', 'first_name', 'username')
    )
    profiles = EmployeeProfile.objects.select_related('department').filter(user__in=users).in_bulk(field_name='user_id')
    employees = []
    for user in users:
        profile = profiles.get(user.id)
        middle_name = profile.middle_name if profile else ''
        full_name = " ".join(
            part for part in [user.last_name, user.first_name, middle_name] if part
        ).strip() or user.username
        department_name = profile.department.name if profile and profile.department else ''
        employees.append(
            {
                'id': user.id,
                'full_name': full_name,
                'email': user.email or '—',
                'role_code': user.role,
                'role_display': user.get_role_display(),
                'department': department_name or '—',
                'department_value': department_name,
                'position': profile.position if profile else '',
                'corporate_phone': profile.corporate_phone if profile else '',
                'monthly_salary': profile.monthly_salary if profile else None,
                'is_active': user.is_active,
                'status_label': 'Активен' if user.is_active else 'Деактивирован',
                'status_class': 'success' if user.is_active else 'warning',
            }
        )

    total_users = len(users)
    active_users = sum(1 for user in users if user.is_active)
    inactive_users = total_users - active_users
    departments = list(Department.objects.order_by('name'))
    reset_requests = list(
        PasswordResetRequest.objects.filter(status='pending')
        .select_related('user')
        .order_by('-created_at')
    )
    reset_request_items = [
        {
            'id': req.id,
            'email': req.email,
            'full_name': req.user.get_full_name() or req.user.username,
            'created_at': req.created_at,
        }
        for req in reset_requests
    ]

    department_positions = _get_department_positions()
    return render(
        request,
        'dashboard/admin/users.html',
        {
            'active_tab': 'users',
            'page_title': 'Пользователи',
            'page_subtitle': 'Управление ролями и доступом',
            'employees': employees,
            'departments': departments,
            'reset_requests': reset_request_items,
            'department_positions': department_positions,
            'stats': {
                'total_users': total_users,
                'active_users': active_users,
                'inactive_users': inactive_users,
                'departments_total': len(departments),
                'pending_requests': len(reset_request_items),
                'last_updated': timezone.localtime().strftime('%H:%M'),
            },
        },
    )


@login_required
@ensure_csrf_cookie
def admin_user_detail(request, user_id):
    _ensure_role(request, 'admin')
    User = get_user_model()
    user = get_object_or_404(User, id=user_id)
    if user.role == 'admin':
        raise PermissionDenied
    profile = getattr(user, 'profile', None)
    department = profile.department.name if profile and profile.department else '—'
    departments = Department.objects.order_by('name')
    department_positions = _get_department_positions()
    avatar_url = ''
    if profile and profile.avatar:
        avatar_url = request.build_absolute_uri(profile.avatar.url)
    employee = {
        'id': user.id,
        'first_name': user.first_name,
        'last_name': user.last_name,
        'middle_name': getattr(profile, 'middle_name', '') if profile else '',
        'full_name': user.get_full_name() or user.username,
        'email': user.email,
        'role': user.role,
        'role_display': user.get_role_display(),
        'department': department,
        'department_id': profile.department_id if profile and profile.department_id else None,
        'position': getattr(profile, 'position', ''),
        'position_display': 'Менеджер' if user.role == 'manager' else getattr(profile, 'position', ''),
        'corporate_phone': getattr(profile, 'corporate_phone', ''),
        'personal_phone': getattr(profile, 'personal_phone', ''),
        'address': getattr(profile, 'address', ''),
        'monthly_salary': getattr(profile, 'monthly_salary', None),
        'salary_reason': getattr(profile, 'salary_reason', ''),
        'passport_series': getattr(profile, 'passport_series', ''),
        'passport_number': getattr(profile, 'passport_number', ''),
        'passport_issued_by': getattr(profile, 'passport_issued_by', ''),
        'passport_issue_date': profile.passport_issue_date.strftime('%d.%m.%Y')
        if profile and profile.passport_issue_date
        else '',
        'snils': getattr(profile, 'snils', ''),
        'inn': getattr(profile, 'inn', ''),
        'avatar_url': avatar_url,
        'status_label': 'Активен' if user.is_active else 'Неактивен',
        'status_class': 'success' if user.is_active else 'warning',
        'is_active': user.is_active,
    }
    back_url = _resolve_back_url(request, reverse("admin-users"))
    return render(
        request,
        'dashboard/admin/user_detail.html',
        {
            'active_tab': 'users',
            'page_title': 'Карточка сотрудника',
            'employee': employee,
            'departments': departments,
            'department_positions': department_positions,
            'back_url': back_url,
        },
    )


@login_required
@ensure_csrf_cookie
def profile_view(request):
    user = request.user
    profile = getattr(user, 'profile', None)
    department = profile.department.name if profile and profile.department else '—'
    avatar_url = ''
    if profile and profile.avatar:
        avatar_url = request.build_absolute_uri(profile.avatar.url)
    departments = Department.objects.order_by('name')
    department_positions = _get_department_positions()

    employee = {
        'id': user.id,
        'first_name': user.first_name,
        'last_name': user.last_name,
        'middle_name': getattr(profile, 'middle_name', '') if profile else '',
        'full_name': user.get_full_name() or user.username,
        'email': user.email,
        'role': user.role,
        'role_display': user.get_role_display(),
        'department': department,
        'department_id': profile.department_id if profile and profile.department_id else None,
        'position': getattr(profile, 'position', ''),
        'position_display': 'Менеджер' if user.role == 'manager' else getattr(profile, 'position', ''),
        'corporate_phone': getattr(profile, 'corporate_phone', ''),
        'personal_phone': getattr(profile, 'personal_phone', ''),
        'address': getattr(profile, 'address', ''),
        'monthly_salary': getattr(profile, 'monthly_salary', None),
        'salary_reason': getattr(profile, 'salary_reason', ''),
        'passport_series': getattr(profile, 'passport_series', ''),
        'passport_number': getattr(profile, 'passport_number', ''),
        'passport_issued_by': getattr(profile, 'passport_issued_by', ''),
        'passport_issue_date': profile.passport_issue_date.strftime('%d.%m.%Y')
        if profile and profile.passport_issue_date
        else '',
        'snils': getattr(profile, 'snils', ''),
        'inn': getattr(profile, 'inn', ''),
        'avatar_url': avatar_url,
        'status_label': 'Активен' if user.is_active else 'Неактивен',
        'status_class': 'success' if user.is_active else 'warning',
        'is_active': user.is_active,
    }

    role = getattr(user, 'role', 'employee')
    base_template = {
        'admin': 'dashboard/admin.html',
        'manager': 'dashboard/manager.html',
        'employee': 'dashboard/employee.html',
    }.get(role, 'dashboard/employee.html')

    back_url_name = {
        'admin': 'admin-dashboard',
        'manager': 'manager-dashboard',
        'employee': 'employee-dashboard',
    }.get(role, 'employee-dashboard')
    back_url = _resolve_back_url(request, reverse(back_url_name))

    return render(
        request,
        'dashboard/profile.html',
        {
            'active_tab': '',
            'page_title': 'Мой профиль',
            'employee': employee,
            'base_template': base_template,
            'back_url': back_url,
            'departments': departments,
            'department_positions': department_positions,
        },
    )


@login_required
def admin_departments(request):
    _ensure_role(request, 'admin')
    User = get_user_model()
    departments = (
        Department.objects.select_related('manager')
        .prefetch_related('positions')
        .annotate(
            employees_count=Count('employees', distinct=True),
            positions_count=Count('positions', distinct=True),
        )
        .order_by('name')
    )
    managers = User.objects.filter(role='manager', is_active=True).order_by(
        'last_name',
        'first_name',
        'username',
    )
    departments_data = [
        {
            'id': department.id,
            'name': department.name,
            'manager_id': department.manager_id,
            'manager_name': department.manager.get_full_name() or department.manager.username
            if department.manager
            else '',
            'positions': list(
                department.positions.order_by('title').values_list('title', flat=True)
            ),
            'positions_count': getattr(department, 'positions_count', 0) or 0,
            'employees_count': getattr(department, 'employees_count', 0) or 0,
            'is_archived': department.is_archived,
        }
        for department in departments
    ]
    total_departments = len(departments_data)
    total_positions = DepartmentPosition.objects.count()
    total_employees = EmployeeProfile.objects.exclude(user__role='admin').count()
    assigned_managers = sum(1 for department in departments_data if department['manager_id'])
    stats = {
        'total_departments': total_departments,
        'total_positions': total_positions,
        'total_employees': total_employees,
        'assigned_managers': assigned_managers,
        'last_updated': timezone.localtime().strftime('%d.%m.%Y %H:%M'),
    }
    return render(
        request,
        'dashboard/admin/departments.html',
        {
            'active_tab': 'departments',
            'page_title': 'Отделы',
            'page_subtitle': 'Правила и лимиты по командам',
            'departments': departments,
            'departments_data': departments_data,
            'managers': managers,
            'stats': stats,
        },
    )


@login_required
def admin_department_detail(request, department_id):
    _ensure_role(request, 'admin')
    department = get_object_or_404(
        Department.objects.select_related('manager'),
        id=department_id,
    )
    positions = list(department.positions.order_by('title'))
    if not positions:
        fallback_positions = (
            EmployeeProfile.objects.filter(department=department)
            .exclude(position='')
            .order_by('position')
            .values_list('position', flat=True)
            .distinct()
        )
        positions = [{'title': title} for title in fallback_positions]
    positions_values = [position['title'] if isinstance(position, dict) else position.title for position in positions]
    User = get_user_model()
    managers = User.objects.filter(role='manager', is_active=True).order_by(
        'last_name',
        'first_name',
        'username',
    )
    users = (
        User.objects.exclude(role__in=['admin', 'manager'])
        .select_related('profile')
        .filter(profile__department=department)
        .order_by('last_name', 'first_name', 'username')
    )
    employees = []
    for user in users:
        profile = getattr(user, 'profile', None)
        employees.append(
            {
                'id': user.id,
                'full_name': user.get_full_name() or user.username,
                'email': user.email,
                'role_display': user.get_role_display(),
                'position': getattr(profile, 'position', '') if profile else '',
                'monthly_salary': getattr(profile, 'monthly_salary', None) if profile else None,
                'is_active': user.is_active,
                'status_label': 'Активен' if user.is_active else 'Неактивен',
                'status_class': 'success' if user.is_active else 'warning',
            }
        )
    available_users = (
        User.objects.exclude(role__in=['admin', 'manager'])
        .exclude(profile__department=department)
        .order_by('last_name', 'first_name', 'username')
    )
    available_employees = [
        {
            'id': user.id,
            'full_name': user.get_full_name() or user.username,
            'email': user.email or '—',
        }
        for user in available_users
    ]
    back_url = _resolve_back_url(request, reverse("admin-departments"))
    return render(
        request,
        'dashboard/admin/department_detail.html',
        {
            'active_tab': 'departments',
            'page_title': department.name,
            'page_subtitle': 'Карточка отдела',
            'department': department,
            'positions': positions,
            'positions_values': positions_values,
            'employees': employees,
            'available_employees': available_employees,
            'managers': managers,
            'transfer_departments': Department.objects.exclude(id=department_id).order_by('name'),
            'back_url': back_url,
        },
    )


@login_required
@ensure_csrf_cookie
def admin_settings(request):
    _ensure_role(request, 'admin')
    settings_obj = GlobalSettings.objects.first()
    if not settings_obj:
        settings_obj = GlobalSettings.objects.create()

    return render(
        request,
        'dashboard/admin/settings.html',
        {
            'active_tab': 'settings',
            'page_title': 'Глобальные настройки',
            'page_subtitle': 'Параметры бренда, поддержки и безопасности',
            'settings_data': {
                'platform_name': settings_obj.platform_name,
                'support_email': settings_obj.support_email,
                'support_phone': settings_obj.support_phone,
                'global_announcement': settings_obj.global_announcement,
                'allow_password_reset_requests': settings_obj.allow_password_reset_requests,
                'backup_retention_days': settings_obj.backup_retention_days,
                'updated_at': settings_obj.updated_at,
            },
        },
    )


def _build_reports_period(request):
    today = timezone.localdate()
    preset_options = [
        ('7d', '7 дней'),
        ('30d', '30 дней'),
        ('90d', '90 дней'),
        ('365d', '365 дней'),
        ('custom', 'Кастомный'),
    ]
    preset_to_days = {
        '7d': 7,
        '30d': 30,
        '90d': 90,
        '365d': 365,
    }

    preset = (request.GET.get('preset') or '30d').strip().lower()
    if preset not in {value for value, _label in preset_options}:
        preset = '30d'

    def parse_date(value):
        text = (value or '').strip()
        if not text:
            return None
        try:
            return date.fromisoformat(text)
        except ValueError:
            return None

    start_input = parse_date(request.GET.get('start_date'))
    end_input = parse_date(request.GET.get('end_date'))

    if preset == 'custom':
        period_end_date = end_input or today
        period_start_date = start_input or (period_end_date - timedelta(days=29))
    else:
        period_end_date = today
        period_start_date = today - timedelta(days=preset_to_days[preset] - 1)

    if period_start_date > period_end_date:
        period_start_date, period_end_date = period_end_date, period_start_date

    max_period_days = 366
    if (period_end_date - period_start_date).days + 1 > max_period_days:
        period_start_date = period_end_date - timedelta(days=max_period_days - 1)
        preset = 'custom'

    period_days = (period_end_date - period_start_date).days + 1
    tz = timezone.get_current_timezone()
    period_start_dt = timezone.make_aware(datetime.combine(period_start_date, time.min), tz)
    period_end_dt_exclusive = timezone.make_aware(
        datetime.combine(period_end_date + timedelta(days=1), time.min),
        tz,
    )

    section = (request.GET.get('section') or 'stats').strip().lower()
    if section not in {'stats', 'reports'}:
        section = 'stats'

    return {
        'preset': preset,
        'preset_options': preset_options,
        'period_start_date': period_start_date,
        'period_end_date': period_end_date,
        'period_start_dt': period_start_dt,
        'period_end_dt_exclusive': period_end_dt_exclusive,
        'period_days': period_days,
        'start_date_value': period_start_date.isoformat(),
        'end_date_value': period_end_date.isoformat(),
        'section': section,
        'print_mode': request.GET.get('print') == '1',
    }


@login_required
def admin_reports(request):
    _ensure_role(request, 'admin')
    maybe_create_daily_backup()
    period_config = _build_reports_period(request)
    period_start_date = period_config['period_start_date']
    period_end_date = period_config['period_end_date']
    period_start_dt = period_config['period_start_dt']
    period_end_dt_exclusive = period_config['period_end_dt_exclusive']
    period_days = int(period_config['period_days'])

    now = timezone.now()

    def pct(part, total, digits=1):
        if not total:
            return 0
        return round((part / total) * 100, digits)

    logs_period_qs = SystemLogEntry.objects.filter(
        created_at__gte=period_start_dt,
        created_at__lt=period_end_dt_exclusive,
    )
    logs_24h_qs = SystemLogEntry.objects.filter(created_at__gte=now - timedelta(hours=24))

    logs_total_period = logs_period_qs.count()
    logs_errors_period = logs_period_qs.filter(level='error').count()
    logs_warnings_period = logs_period_qs.filter(level='warning').count()
    logs_slow_period = logs_period_qs.filter(duration_ms__gte=1000).count()

    logs_total_24h = logs_24h_qs.count()
    logs_errors_24h = logs_24h_qs.filter(level='error').count()
    logs_warnings_24h = logs_24h_qs.filter(level='warning').count()
    logs_slow_24h = logs_24h_qs.filter(duration_ms__gte=1000).count()

    avg_response_24h = logs_24h_qs.exclude(duration_ms__isnull=True).aggregate(avg_duration=Avg('duration_ms')).get(
        'avg_duration'
    )
    avg_response_period = logs_period_qs.exclude(duration_ms__isnull=True).aggregate(avg_duration=Avg('duration_ms')).get(
        'avg_duration'
    )
    max_response_period = logs_period_qs.exclude(duration_ms__isnull=True).aggregate(max_duration=Max('duration_ms')).get(
        'max_duration'
    )

    avg_response_24h_ms = int(round(avg_response_24h)) if avg_response_24h is not None else None
    avg_response_period_ms = int(round(avg_response_period)) if avg_response_period is not None else None
    max_response_period_ms = int(round(max_response_period)) if max_response_period is not None else None

    online_users = (
        SystemLogEntry.objects.filter(created_at__gte=now - timedelta(minutes=5), user__isnull=False)
        .values('user_id')
        .distinct()
        .count()
    )
    requests_per_minute = SystemLogEntry.objects.filter(created_at__gte=now - timedelta(minutes=1)).count()
    active_users_period = logs_period_qs.filter(user__isnull=False).values('user_id').distinct().count()

    status_base_qs = logs_period_qs.exclude(status_code__isnull=True)
    status_total = status_base_qs.count()
    status_2xx = status_base_qs.filter(status_code__gte=200, status_code__lt=300).count()
    status_3xx = status_base_qs.filter(status_code__gte=300, status_code__lt=400).count()
    status_4xx = status_base_qs.filter(status_code__gte=400, status_code__lt=500).count()
    status_5xx = status_base_qs.filter(status_code__gte=500).count()
    status_other = max(status_total - status_2xx - status_3xx - status_4xx - status_5xx, 0)
    status_5xx_rate = pct(status_5xx, status_total)

    status_distribution = [
        {'code': 's2xx', 'label': '2xx', 'value': status_2xx, 'share': pct(status_2xx, status_total)},
        {'code': 's3xx', 'label': '3xx', 'value': status_3xx, 'share': pct(status_3xx, status_total)},
        {'code': 's4xx', 'label': '4xx', 'value': status_4xx, 'share': pct(status_4xx, status_total)},
        {'code': 's5xx', 'label': '5xx', 'value': status_5xx, 'share': pct(status_5xx, status_total)},
        {'code': 'other', 'label': 'Other', 'value': status_other, 'share': pct(status_other, status_total)},
    ]

    daily_requests_map = {
        row['day']: row['total']
        for row in (
            logs_period_qs.annotate(day=TruncDate('created_at'))
            .values('day')
            .annotate(total=Count('id'))
        )
    }
    daily_errors_map = {
        row['day']: row['total']
        for row in (
            logs_period_qs.filter(level='error')
            .annotate(day=TruncDate('created_at'))
            .values('day')
            .annotate(total=Count('id'))
        )
    }
    daily_warnings_map = {
        row['day']: row['total']
        for row in (
            logs_period_qs.filter(level='warning')
            .annotate(day=TruncDate('created_at'))
            .values('day')
            .annotate(total=Count('id'))
        )
    }
    daily_avg_duration_map = {
        row['day']: int(round(row['avg_duration'] or 0))
        for row in (
            logs_period_qs.exclude(duration_ms__isnull=True)
            .annotate(day=TruncDate('created_at'))
            .values('day')
            .annotate(avg_duration=Avg('duration_ms'))
        )
    }
    daily_max_duration_map = {
        row['day']: int(round(row['max_duration'] or 0))
        for row in (
            logs_period_qs.exclude(duration_ms__isnull=True)
            .annotate(day=TruncDate('created_at'))
            .values('day')
            .annotate(max_duration=Max('duration_ms'))
        )
    }

    def parse_percent(value):
        try:
            number = float(value)
        except (TypeError, ValueError):
            return None
        if number < 0:
            return None
        return number

    resource_daily_samples = {}
    for row in logs_period_qs.values('created_at', 'meta'):
        raw_meta = row.get('meta')
        if not isinstance(raw_meta, dict):
            continue
        day = timezone.localtime(row['created_at']).date()
        bucket = resource_daily_samples.setdefault(day, {'cpu': [], 'memory': [], 'disk': []})
        cpu = parse_percent(raw_meta.get('cpu_load_percent_1m'))
        memory = parse_percent(raw_meta.get('memory_used_percent'))
        disk = parse_percent(raw_meta.get('disk_used_percent'))
        if cpu is not None:
            bucket['cpu'].append(cpu)
        if memory is not None:
            bucket['memory'].append(memory)
        if disk is not None:
            bucket['disk'].append(disk)

    chart_labels = []
    chart_requests = []
    chart_errors = []
    chart_warnings = []
    chart_avg_ms = []
    chart_max_ms = []
    chart_cpu = []
    chart_memory = []
    chart_disk = []
    max_traffic_value = 0
    max_latency_value = 0
    max_resource_value = 0

    current_day = period_start_date
    while current_day <= period_end_date:
        requests_count = daily_requests_map.get(current_day, 0)
        errors_count = daily_errors_map.get(current_day, 0)
        warnings_count = daily_warnings_map.get(current_day, 0)
        avg_ms = daily_avg_duration_map.get(current_day, 0)
        max_ms = daily_max_duration_map.get(current_day, 0)
        resource_bucket = resource_daily_samples.get(current_day, {})
        cpu_values = resource_bucket.get('cpu') or []
        memory_values = resource_bucket.get('memory') or []
        disk_values = resource_bucket.get('disk') or []
        cpu_avg = round(sum(cpu_values) / len(cpu_values), 1) if cpu_values else 0
        memory_avg = round(sum(memory_values) / len(memory_values), 1) if memory_values else 0
        disk_avg = round(sum(disk_values) / len(disk_values), 1) if disk_values else 0

        chart_labels.append(current_day.strftime('%d.%m'))
        chart_requests.append(requests_count)
        chart_errors.append(errors_count)
        chart_warnings.append(warnings_count)
        chart_avg_ms.append(avg_ms)
        chart_max_ms.append(max_ms)
        chart_cpu.append(cpu_avg)
        chart_memory.append(memory_avg)
        chart_disk.append(disk_avg)

        max_traffic_value = max(max_traffic_value, requests_count, errors_count, warnings_count)
        max_latency_value = max(max_latency_value, avg_ms, max_ms)
        max_resource_value = max(max_resource_value, cpu_avg, memory_avg, disk_avg)

        current_day += timedelta(days=1)

    traffic_chart = {
        'labels': chart_labels,
        'series': {
            'requests': chart_requests,
            'errors': chart_errors,
            'warnings': chart_warnings,
        },
        'max_value': max_traffic_value,
    }
    latency_chart = {
        'labels': chart_labels,
        'series': {
            'avg_ms': chart_avg_ms,
            'max_ms': chart_max_ms,
        },
        'max_value': max_latency_value,
    }
    resource_chart = {
        'labels': chart_labels,
        'series': {
            'cpu': chart_cpu,
            'memory': chart_memory,
            'disk': chart_disk,
        },
        'max_value': max_resource_value,
    }

    top_endpoints = []
    top_endpoints_qs = (
        logs_period_qs.exclude(path='')
        .values('method', 'path')
        .annotate(
            total=Count('id'),
            errors=Count('id', filter=Q(level='error')),
            avg_duration=Avg('duration_ms'),
            max_duration=Max('duration_ms'),
            last_seen=Max('created_at'),
        )
        .order_by('-total', 'path')[:12]
    )
    for endpoint in top_endpoints_qs:
        method = (endpoint.get('method') or '—').upper()
        path = endpoint.get('path') or '/'
        endpoint_label = f'{method} {path}'.strip()
        if len(endpoint_label) > 88:
            endpoint_label = f'{endpoint_label[:85]}...'
        avg_duration = endpoint.get('avg_duration')
        max_duration = endpoint.get('max_duration')
        top_endpoints.append({
            'method': method,
            'path': path,
            'label': endpoint_label,
            'total': endpoint['total'],
            'errors': endpoint.get('errors', 0),
            'error_rate': pct(endpoint.get('errors', 0), endpoint['total']),
            'avg_duration_ms': int(round(avg_duration)) if avg_duration is not None else None,
            'max_duration_ms': int(round(max_duration)) if max_duration is not None else None,
            'last_seen': timezone.localtime(endpoint['last_seen']) if endpoint.get('last_seen') else None,
        })

    last_backup = SystemBackup.objects.filter(status='ready').order_by('-created_at').first()
    backup_qs = SystemBackup.objects.filter(
        created_at__gte=period_start_dt,
        created_at__lt=period_end_dt_exclusive,
    )
    backups_created_period = backup_qs.count()
    backup_size_period = backup_qs.aggregate(total_size=Sum('file_size')).get('total_size') or 0
    failed_backups_period = backup_qs.filter(status='failed').count()
    restored_backups_period = backup_qs.filter(status='restored').count()

    last_backup_age_hours = None
    if last_backup:
        last_backup_age_hours = int((now - last_backup.created_at).total_seconds() // 3600)

    runtime_metrics = collect_runtime_metrics()
    cpu_load_percent_1m = runtime_metrics.get('cpu_load_percent_1m')
    memory_used_percent = runtime_metrics.get('memory_used_percent')
    memory_used_bytes = runtime_metrics.get('memory_used_bytes')
    memory_total_bytes = runtime_metrics.get('memory_total_bytes')
    memory_source = runtime_metrics.get('memory_source') or 'unknown'
    disk_used_percent = runtime_metrics.get('disk_used_percent')
    disk_used_bytes = runtime_metrics.get('disk_used_bytes')
    disk_total_bytes = runtime_metrics.get('disk_total_bytes')

    def load_tone(value, warning_level=80, danger_level=92):
        if value is None:
            return 'muted'
        if value >= danger_level:
            return 'danger'
        if value >= warning_level:
            return 'warning'
        return 'success'

    resource_cards = [
        {
            'code': 'cpu',
            'label': 'CPU load (1m)',
            'value': cpu_load_percent_1m,
            'value_label': f'{cpu_load_percent_1m:.1f}%' if cpu_load_percent_1m is not None else '—',
            'tone': load_tone(cpu_load_percent_1m, warning_level=75, danger_level=90),
            'meta': (
                f"Ядер: {runtime_metrics.get('cpu_count', 0)} · load: "
                f"{runtime_metrics.get('load_avg_1m'):.2f}"
                if runtime_metrics.get('load_avg_1m') is not None
                else 'Нагрузка CPU недоступна'
            ),
        },
        {
            'code': 'memory',
            'label': 'Память',
            'value': memory_used_percent,
            'value_label': f'{memory_used_percent:.1f}%' if memory_used_percent is not None else '—',
            'tone': load_tone(memory_used_percent, warning_level=78, danger_level=90),
            'meta': (
                f"{format_bytes(memory_used_bytes)} / {format_bytes(memory_total_bytes)}"
                if memory_used_bytes is not None and memory_total_bytes is not None
                else f"{format_bytes(memory_used_bytes)} (процесс)"
                if memory_used_bytes is not None
                else 'Память недоступна'
            ),
            'source': memory_source,
        },
        {
            'code': 'disk',
            'label': 'Диск',
            'value': disk_used_percent,
            'value_label': f'{disk_used_percent:.1f}%' if disk_used_percent is not None else '—',
            'tone': load_tone(disk_used_percent, warning_level=82, danger_level=92),
            'meta': (
                f"{format_bytes(disk_used_bytes)} / {format_bytes(disk_total_bytes)}"
                if disk_used_bytes is not None and disk_total_bytes is not None
                else 'Диск недоступен'
            ),
        },
    ]

    critical_issues = []

    if cpu_load_percent_1m is not None:
        if cpu_load_percent_1m >= 95:
            critical_issues.append({
                'tone': 'danger',
                'title': 'CPU перегружен',
                'details': f'Текущая нагрузка CPU {cpu_load_percent_1m:.1f}%.',
            })
        elif cpu_load_percent_1m >= 80:
            critical_issues.append({
                'tone': 'warning',
                'title': 'Высокая нагрузка CPU',
                'details': f'Текущая нагрузка CPU {cpu_load_percent_1m:.1f}%.',
            })

    if memory_used_percent is not None:
        if memory_used_percent >= 95:
            critical_issues.append({
                'tone': 'danger',
                'title': 'Память почти исчерпана',
                'details': f'Использование памяти {memory_used_percent:.1f}%.',
            })
        elif memory_used_percent >= 85:
            critical_issues.append({
                'tone': 'warning',
                'title': 'Высокая загрузка памяти',
                'details': f'Использование памяти {memory_used_percent:.1f}%.',
            })

    if disk_used_percent is not None:
        if disk_used_percent >= 95:
            critical_issues.append({
                'tone': 'danger',
                'title': 'Мало места на диске',
                'details': f'Использование диска {disk_used_percent:.1f}%.',
            })
        elif disk_used_percent >= 85:
            critical_issues.append({
                'tone': 'warning',
                'title': 'Диск близок к заполнению',
                'details': f'Использование диска {disk_used_percent:.1f}%.',
            })

    if logs_errors_24h >= 50:
        critical_issues.append({
            'tone': 'danger',
            'title': 'Много ошибок за 24 часа',
            'details': f'За последние сутки зарегистрировано {logs_errors_24h} ошибок.',
        })
    elif logs_errors_24h >= 10:
        critical_issues.append({
            'tone': 'warning',
            'title': 'Растет число ошибок',
            'details': f'За последние сутки зарегистрировано {logs_errors_24h} ошибок.',
        })

    if avg_response_24h_ms is not None:
        if avg_response_24h_ms >= 1500:
            critical_issues.append({
                'tone': 'danger',
                'title': 'Высокая задержка ответа',
                'details': f'Среднее время ответа за 24ч: {avg_response_24h_ms} мс.',
            })
        elif avg_response_24h_ms >= 800:
            critical_issues.append({
                'tone': 'warning',
                'title': 'Задержка ответа выше нормы',
                'details': f'Среднее время ответа за 24ч: {avg_response_24h_ms} мс.',
            })

    if status_5xx_rate >= 3:
        critical_issues.append({
            'tone': 'danger',
            'title': 'Высокая доля 5xx',
            'details': f'Доля ответов 5xx за период: {status_5xx_rate:.1f}%.',
        })
    elif status_5xx_rate >= 1:
        critical_issues.append({
            'tone': 'warning',
            'title': 'Нестабильные ответы 5xx',
            'details': f'Доля ответов 5xx за период: {status_5xx_rate:.1f}%.',
        })

    if last_backup_age_hours is None:
        critical_issues.append({
            'tone': 'danger',
            'title': 'Нет готовых бэкапов',
            'details': 'Система не обнаружила ни одной актуальной резервной копии.',
        })
    elif last_backup_age_hours > 48:
        critical_issues.append({
            'tone': 'warning',
            'title': 'Бэкап давно не обновлялся',
            'details': f'Последний готовый бэкап создан {last_backup_age_hours} ч назад.',
        })

    if failed_backups_period > 0:
        critical_issues.append({
            'tone': 'warning',
            'title': 'Есть неуспешные бэкапы',
            'details': f'За период не удалось создать {failed_backups_period} бэкапов.',
        })

    if not critical_issues:
        critical_issues.append({
            'tone': 'success',
            'title': 'Критических проблем не обнаружено',
            'details': 'Нагрузка и ошибки находятся в допустимых пределах.',
        })

    health_score = 100
    for issue in critical_issues:
        if issue['tone'] == 'danger':
            health_score -= 25
        elif issue['tone'] == 'warning':
            health_score -= 12
    health_score = max(0, min(100, health_score))

    if health_score >= 85:
        health_tone = 'success'
        health_label = 'Стабильно'
    elif health_score >= 65:
        health_tone = 'warning'
        health_label = 'Требует внимания'
    else:
        health_tone = 'danger'
        health_label = 'Есть риски'

    stats = {
        'period_days': period_days,
        'period_start': period_start_date,
        'period_end': period_end_date,
        'period_label': f'{period_start_date:%d.%m.%Y} — {period_end_date:%d.%m.%Y}',
        'logs_total_period': logs_total_period,
        'logs_errors_period': logs_errors_period,
        'logs_warnings_period': logs_warnings_period,
        'logs_slow_period': logs_slow_period,
        'logs_total_24h': logs_total_24h,
        'logs_errors_24h': logs_errors_24h,
        'logs_warnings_24h': logs_warnings_24h,
        'logs_slow_24h': logs_slow_24h,
        'avg_response_24h_ms': avg_response_24h_ms,
        'avg_response_period_ms': avg_response_period_ms,
        'max_response_period_ms': max_response_period_ms,
        'status_5xx_rate': status_5xx_rate,
        'status_5xx': status_5xx,
        'requests_per_minute': requests_per_minute,
        'online_users': online_users,
        'active_users_period': active_users_period,
        'cpu_load_percent_1m': cpu_load_percent_1m,
        'memory_used_percent': memory_used_percent,
        'memory_used_bytes': memory_used_bytes,
        'memory_total_bytes': memory_total_bytes,
        'memory_source': memory_source,
        'disk_used_percent': disk_used_percent,
        'disk_used_bytes': disk_used_bytes,
        'disk_total_bytes': disk_total_bytes,
        'backups_created_period': backups_created_period,
        'backup_size_period': backup_size_period,
        'failed_backups_period': failed_backups_period,
        'restored_backups_period': restored_backups_period,
        'last_backup': last_backup,
        'last_backup_age_hours': last_backup_age_hours,
        'health_score': health_score,
        'health_label': health_label,
        'health_tone': health_tone,
        'last_updated': timezone.localtime(now).strftime('%d.%m.%Y %H:%M'),
    }

    return render(
        request,
        'dashboard/admin/reports.html',
        {
            'active_tab': 'reports',
            'page_title': 'Мониторинг сайта',
            'page_subtitle': 'Загрузка, ошибки, производительность и устойчивость платформы',
            'stats': stats,
            'resource_cards': resource_cards,
            'critical_issues': critical_issues,
            'top_endpoints': top_endpoints,
            'status_distribution': status_distribution,
            'traffic_chart': traffic_chart,
            'latency_chart': latency_chart,
            'resource_chart': resource_chart,
            'filters': {
                'preset': period_config['preset'],
                'preset_options': period_config['preset_options'],
                'start_date': period_config['start_date_value'],
                'end_date': period_config['end_date_value'],
                'period_days': period_days,
            },
            'print_mode': period_config['print_mode'],
        },
    )


@login_required
def admin_reports_export(request):
    _ensure_role(request, 'admin')
    export_format = (request.GET.get('format') or 'xlsx').lower()
    period_config = _build_reports_period(request)
    period_start_date = period_config['period_start_date']
    period_end_date = period_config['period_end_date']
    period_start_dt = period_config['period_start_dt']
    period_end_dt_exclusive = period_config['period_end_dt_exclusive']

    User = get_user_model()
    now = timezone.now()
    today = timezone.localdate()

    departments = list(
        Department.objects.select_related('manager')
        .annotate(
            employees_count=Count('employees', filter=Q(employees__user__role='employee'), distinct=True),
            positions_count=Count('positions', distinct=True),
        )
        .order_by('name')
    )
    open_tasks_by_department = {
        row['department_id']: row['total']
        for row in (
            DepartmentTask.objects.filter(status__in=['todo', 'in_progress'])
            .values('department_id')
            .annotate(total=Count('id'))
        )
    }
    overdue_tasks_by_department = {
        row['department_id']: row['total']
        for row in (
            DepartmentTask.objects.filter(status__in=['todo', 'in_progress'], date__lt=today)
            .values('department_id')
            .annotate(total=Count('id'))
        )
    }
    tasks_created_period_by_department = {
        row['department_id']: row['total']
        for row in (
            DepartmentTask.objects.filter(
                created_at__gte=period_start_dt,
                created_at__lt=period_end_dt_exclusive,
            )
            .values('department_id')
            .annotate(total=Count('id'))
        )
    }
    pending_requests_by_department = {
        row['user__profile__department_id']: row['total']
        for row in (
            EmployeeShiftRequest.objects.filter(status='pending', user__profile__department__isnull=False)
            .values('user__profile__department_id')
            .annotate(total=Count('id'))
        )
    }
    requests_created_period_by_department = {
        row['user__profile__department_id']: row['total']
        for row in (
            EmployeeShiftRequest.objects.filter(
                created_at__gte=period_start_dt,
                created_at__lt=period_end_dt_exclusive,
                user__profile__department__isnull=False,
            )
            .values('user__profile__department_id')
            .annotate(total=Count('id'))
        )
    }

    department_rows = []
    for department in departments:
        manager_name = '—'
        if department.manager:
            manager_name = department.manager.get_full_name() or department.manager.username
        department_rows.append(
            [
                department.name,
                'Архивный' if department.is_archived else 'Активный',
                manager_name,
                department.employees_count,
                department.positions_count,
                open_tasks_by_department.get(department.id, 0),
                overdue_tasks_by_department.get(department.id, 0),
                pending_requests_by_department.get(department.id, 0),
                tasks_created_period_by_department.get(department.id, 0),
                requests_created_period_by_department.get(department.id, 0),
            ]
        )

    if export_format == 'csv':
        output = io.StringIO()
        writer = csv.writer(output)
        output.write('\ufeff')
        writer.writerow(
            [
                'Отдел',
                'Статус',
                'Менеджер',
                'Сотрудников',
                'Должностей',
                'Открытых задач',
                'Просроченных задач',
                'Запросов на рассмотрении',
                'Новых задач за период',
                'Новых запросов за период',
            ]
        )
        for row in department_rows:
            writer.writerow(row)
        response = HttpResponse(output.getvalue(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = f'attachment; filename="reports_departments_{today:%Y%m%d}.csv"'
        return response

    tasks_created_period = DepartmentTask.objects.filter(
        created_at__gte=period_start_dt,
        created_at__lt=period_end_dt_exclusive,
    ).count()
    tasks_done_period = DepartmentTask.objects.filter(
        status='done',
        updated_at__gte=period_start_dt,
        updated_at__lt=period_end_dt_exclusive,
    ).count()
    requests_created_period = EmployeeShiftRequest.objects.filter(
        created_at__gte=period_start_dt,
        created_at__lt=period_end_dt_exclusive,
    ).count()

    summary_rows = [
        ('Период отчета', f'{period_start_date:%d.%m.%Y} — {period_end_date:%d.%m.%Y} ({period_config["period_days"]} дн.)'),
        ('Дата формирования', timezone.localtime(now).strftime('%d.%m.%Y %H:%M')),
        ('Пользователей (без админов)', User.objects.exclude(role='admin').count()),
        ('Активных пользователей', User.objects.exclude(role='admin').filter(is_active=True).count()),
        ('Отделов активных', Department.objects.filter(is_archived=False).count()),
        ('Задач открытых', DepartmentTask.objects.filter(status__in=['todo', 'in_progress']).count()),
        ('Задач просроченных', DepartmentTask.objects.filter(status__in=['todo', 'in_progress'], date__lt=today).count()),
        ('Задач создано за период', tasks_created_period),
        ('Задач закрыто за период', tasks_done_period),
        ('Запросов создано за период', requests_created_period),
        ('Запросов на рассмотрении', EmployeeShiftRequest.objects.filter(status='pending').count()),
        ('Ошибок в логах за 24ч', SystemLogEntry.objects.filter(created_at__gte=now - timedelta(hours=24), level='error').count()),
        (
            'Неудачных бэкапов за период',
            SystemBackup.objects.filter(
                status='failed',
                created_at__gte=period_start_dt,
                created_at__lt=period_end_dt_exclusive,
            ).count(),
        ),
    ]

    top_endpoints = list(
        SystemLogEntry.objects.filter(created_at__gte=now - timedelta(hours=24))
        .exclude(path='')
        .values('method', 'path')
        .annotate(total=Count('id'))
        .order_by('-total', 'path')[:20]
    )

    workbook = Workbook()
    sheet_summary = workbook.active
    sheet_summary.title = 'Сводка'
    sheet_summary.append(['Показатель', 'Значение'])
    for row in summary_rows:
        sheet_summary.append(list(row))

    header_fill = PatternFill(fill_type='solid', fgColor='1F2937')
    header_font = Font(bold=True, color='FFFFFF')

    for cell in sheet_summary[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet_summary.column_dimensions['A'].width = 46
    sheet_summary.column_dimensions['B'].width = 24
    sheet_summary.freeze_panes = 'A2'
    sheet_summary.auto_filter.ref = sheet_summary.dimensions

    sheet_departments = workbook.create_sheet('Отделы')
    department_headers = [
        'Отдел',
        'Статус',
        'Менеджер',
        'Сотрудников',
        'Должностей',
        'Открытых задач',
        'Просроченных задач',
        'Запросов pending',
        'Новых задач (период)',
        'Новых запросов (период)',
    ]
    sheet_departments.append(department_headers)
    for row in department_rows:
        sheet_departments.append(row)
    for cell in sheet_departments[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    department_widths = [28, 14, 28, 14, 14, 16, 19, 18, 20, 20]
    for index, width in enumerate(department_widths, start=1):
        column_name = get_column_letter(index)
        sheet_departments.column_dimensions[column_name].width = width
    sheet_departments.freeze_panes = 'A2'
    sheet_departments.auto_filter.ref = sheet_departments.dimensions

    sheet_api = workbook.create_sheet('API 24ч')
    sheet_api.append(['Метод и путь', 'Запросов за 24ч'])
    if top_endpoints:
        for endpoint in top_endpoints:
            method = (endpoint.get('method') or '—').upper()
            path = endpoint.get('path') or '/'
            sheet_api.append([f'{method} {path}'.strip(), endpoint['total']])
    else:
        sheet_api.append(['Нет данных', 0])
    for cell in sheet_api[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    sheet_api.column_dimensions['A'].width = 90
    sheet_api.column_dimensions['B'].width = 18
    sheet_api.freeze_panes = 'A2'
    sheet_api.auto_filter.ref = sheet_api.dimensions

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="reports_full_{today:%Y%m%d}.xlsx"'
    return response


@login_required
def admin_payroll(request):
    return _render_admin_page(
        request,
        'dashboard/admin/payroll.html',
        'payroll',
        'Отчеты по времени',
        'Сводка доступности, нагрузки и переработок',
    )


@login_required
@ensure_csrf_cookie
def admin_system(request):
    _ensure_role(request, 'admin')
    maybe_create_daily_backup()
    monitoring = _system_monitoring_snapshot()

    logs = list(SystemLogEntry.objects.select_related('user').order_by('-created_at')[:12])
    log_entries = []
    for entry in logs:
        user_name = "Гость"
        if entry.user:
            user_name = entry.user.get_full_name() or entry.user.username
        log_entries.append(
            {
                "id": entry.id,
                "created_at": timezone.localtime(entry.created_at),
                "action": entry.action,
                "level": entry.level,
                "level_label": entry.get_level_display(),
                "status_code": entry.status_code,
                "method": entry.method,
                "path": entry.path,
                "user_name": user_name,
            }
        )

    backups = list(SystemBackup.objects.select_related('created_by', 'restored_by').order_by('-created_at')[:10])
    for backup in backups:
        backup.size_display = format_bytes(backup.file_size)

    return render(
        request,
        'dashboard/admin/system.html',
        {
            'active_tab': 'system',
            'page_title': 'Система',
            'page_subtitle': 'Логи, бэкапы, мониторинг',
            'system_logs': log_entries,
            'backups': backups,
            'monitoring': monitoring,
            'last_backup': monitoring.get('last_backup'),
        },
    )


@login_required
@require_http_methods(["GET"])
def admin_system_backup_download(request, backup_id):
    _ensure_role(request, 'admin')
    backup = get_object_or_404(SystemBackup, id=backup_id)
    file_path = Path(backup.file_path)
    if not file_path.exists():
        raise Http404("Backup file not found.")
    return FileResponse(
        open(file_path, "rb"),
        as_attachment=True,
        filename=backup.file_name,
    )


@login_required
@require_http_methods(["GET"])
def admin_system_logs_export(request):
    _ensure_role(request, 'admin')
    export_format = (request.GET.get("format") or "xlsx").lower()
    logs = SystemLogEntry.objects.select_related("user").order_by("-created_at")[:2000]
    if export_format == "json":
        return JsonResponse({"logs": [_serialize_log_entry(entry) for entry in logs]})

    headers = ["Время", "Пользователь", "Действие", "Метод", "Путь", "Статус", "Уровень", "IP", "Длительность (мс)"]
    rows = []
    for entry in logs:
        user_name = "Гость"
        if entry.user:
            user_name = entry.user.get_full_name() or entry.user.username
        rows.append(
            [
                timezone.localtime(entry.created_at).strftime("%Y-%m-%d %H:%M:%S"),
                user_name,
                entry.action,
                entry.method,
                entry.path,
                entry.status_code or "",
                entry.get_level_display(),
                entry.ip_address or "",
                entry.duration_ms or "",
            ]
        )

    if export_format == "xlsx":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "System Logs"
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)

        header_fill = PatternFill(fill_type="solid", fgColor="1F2937")
        header_font = Font(bold=True, color="FFFFFF")
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        column_widths = [19, 28, 42, 10, 48, 10, 14, 18, 18]
        for index, width in enumerate(column_widths, start=1):
            worksheet.column_dimensions[chr(64 + index)].width = width

        for row in worksheet.iter_rows(min_row=2, max_col=9):
            row[2].alignment = Alignment(vertical="top", wrap_text=True)
            row[4].alignment = Alignment(vertical="top", wrap_text=True)

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        filename = f"system_logs_{timezone.localdate():%Y%m%d}.xlsx"
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response

    output = io.StringIO()
    writer = csv.writer(output)
    output.write("\ufeff")
    writer.writerow(headers)
    for row in rows:
        writer.writerow(row)
    filename = f"system_logs_{timezone.localdate():%Y%m%d}.csv"
    response = HttpResponse(output.getvalue(), content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
